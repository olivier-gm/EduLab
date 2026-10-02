# -*- coding: utf-8 -*-
"""Subtítulos copiados por la IA tal como se escribieron ('que es...'), líneas
sueltas de puntos y página en blanco antes de la conclusión (caso 'Craneo')."""
import docx
import pytest

from algorythms import Document_process as D
from form_processor import FormProcessor


def build(body, known=(), detect=True):
    document = docx.Document()
    D.parrafos(body, document, 'Craneo', True, detect, known)
    D.parrafos('Final.', document, 'Conclusión', False, detect)
    return document


def test_subtitulo_con_puntos_suspensivos_y_minuscula_queda_como_subtitulo_formal():
    document = build('Intro del tema.\n\n\nque es...\n\n\nUn texto largo.', ['Que es'])
    sub = [p for p in document.paragraphs if p.style.name == 'Heading 2']
    assert [p.text for p in sub] == ['Que es']


def test_puntos_suspensivos_sin_lista_conocida_tambien_son_subtitulo():
    document = build('Intro del tema.\n\n\nhuesos del craneo...\n\n\nUn texto largo.')
    assert [p.text for p in document.paragraphs if p.style.name == 'Heading 2'] == ['Huesos del craneo']


def test_linea_de_solo_puntos_no_genera_parrafo():
    document = build('Uno es el texto.\n\n\n.\n\n\n...\n\n\nDos es el texto.')
    assert '.' not in [p.text.strip() for p in document.paragraphs]
    assert len([p for p in document.paragraphs if p.style.name == 'Normal']) == 3


def test_el_salto_antes_de_la_conclusion_no_es_un_parrafo_suelto():
    document = build('Texto.')
    conclusion = [p for p in document.paragraphs if p.text == 'Conclusión'][0]
    assert conclusion.paragraph_format.page_break_before is True
    assert 'w:br' not in document.element.xml and 'type="page"' not in document.element.xml


@pytest.mark.parametrize('raw,expected', [('que es...', 'Que es'), ('  huesos   del cráneo… ', 'Huesos del cráneo'),
                                           ('nervios:', 'Nervios'), ('Ya bien', 'Ya bien')])
def test_limpia_los_subtitulos_del_usuario(raw, expected):
    assert FormProcessor.clean_subtitle(raw) == expected
