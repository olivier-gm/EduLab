import json
from types import SimpleNamespace

import pytest
from docx import Document

import IA
import ai_provider
import db
import workshop
from algorythms import Document_process
from form_processor import FormProcessor


@pytest.mark.parametrize('form', [{}, {'subtitle_1': '  ...  '}, {'subtitle_1': 'A' * 301},
                                  {'subtitle_1': 'Pregunta', 'subtitle_9': 'Otra'}])
def test_invalid_questions_are_rejected(form):
    with pytest.raises(ValueError):
        workshop.questions_from_form(form)


def test_questions_preserve_order_and_ignore_empty_slots():
    assert workshop.questions_from_form({'subtitle_3': ' ¿Qué es una célula? ',
                                         'subtitle_1': ' Defina la mitosis: ', 'subtitle_2': ''}) == [
        'Defina la mitosis', '¿Qué es una célula?']


@pytest.mark.parametrize('payload', [
    '', '[]', '{}', '{"answers":null}', '{"answers":[]}',
    '{"answers":[{"number":2,"answer":"Texto"}]}',
    '{"answers":[{"number":true,"answer":"Texto"}]}',
    '{"answers":[{"number":1,"answer":"  "}]}',
    '{"answers":[{"number":1,"answer":123}]}',
])
def test_incomplete_or_malformed_response_is_not_delivered(monkeypatch, payload):
    monkeypatch.setattr(ai_provider, 'generate_content', lambda **kw: SimpleNamespace(text=payload))
    with pytest.raises(IA.GenerationError):
        workshop.generate_workshop(['¿Qué es una célula?'])


def test_prompt_uses_only_questions_and_accounts_for_tokens(monkeypatch):
    calls = []
    def generate(**kw):
        calls.append(kw)
        return SimpleNamespace(text=json.dumps({'answers': [
            {'number': 1, 'answer': 'Primera respuesta.\n\nOtro párrafo.'},
            {'number': 2, 'answer': 'Segunda respuesta.'}]}),
            usage_metadata=SimpleNamespace(total_token_count=42))
    monkeypatch.setattr(ai_provider, 'generate_content', generate)
    questions = ['¿Qué es una célula?', 'Explique la mitosis']
    usage = []
    entries = workshop.generate_workshop(questions, usage)
    assert [entry['question'] for entry in entries] == questions
    assert entries[0]['answer'] == 'Primera respuesta.\n\nOtro párrafo.'
    assert usage == [42]
    assert calls[0]['reasoning_purpose'] == 'workshop'
    assert json.loads(calls[0]['contents']) == [
        {'number': 1, 'question': questions[0]}, {'number': 2, 'question': questions[1]}]
    assert IA.EDUCATIONAL_CONTEXT in calls[0]['config'].system_instruction
    assert IA.MATH_INSTRUCTION in calls[0]['config'].system_instruction
    assert calls[0]['config'].response_json_schema['properties']['answers']['type'] == 'array'


@pytest.mark.parametrize('route', ['uni', 'bach'])
@pytest.mark.parametrize('count', [1, 8])
def test_workshop_document_has_cover_then_questions_without_toc(tmp_path, monkeypatch, route, count):
    title = 'Ensayo de biología'  # El título no convierte el taller en un ensayo.
    processor = FormProcessor({'title': title, 'u': 'Institución de prueba'}, route)
    processor.process()
    replacements, head = processor.generate_replacements()
    entries = [{'question': f'Explique el proceso {i}', 'answer': 'Primer párrafo.\n\nSegundo párrafo.'}
               for i in range(1, count + 1)]
    monkeypatch.setattr(Document_process, 'convert', lambda *a: None)
    output = tmp_path / 'taller.docx'
    template = 'input/plantilla_bach.docx' if route == 'bach' else 'input/plantilla.docx'
    Document_process.fill_placeholders(str(output), template, '', replacements,
        '', '', '', head, route, workshop_entries=entries)
    document = Document(output)
    headings = [p for p in document.paragraphs if p.style.name == 'Heading 1']
    assert [p.text for p in headings] == [f'{i}. Explique el proceso {i}' for i in range(1, count + 1)]
    first_heading = next(
        i for i, p in enumerate(document.paragraphs) if p.text == headings[0].text)
    assert document.paragraphs[first_heading - 1]._p.xpath('.//w:br[@w:type="page"]')
    assert all(p.paragraph_format.keep_with_next for p in headings)
    assert sum(p.text == 'Primer párrafo.' for p in document.paragraphs) == count
    assert sum(p.text == 'Segundo párrafo.' for p in document.paragraphs) == count
    assert not document.element.xpath('//w:instrText[contains(text(), "TOC")]')
    assert not any(p.text in ('Introducción', 'Conclusión', 'Bibliografía', 'Índice') for p in document.paragraphs)


@pytest.mark.parametrize('purpose,light,expected', [('content', False, 'low'), ('workshop', False, 'xhigh'),
                                                   ('workshop', True, 'low')])
def test_reasoning_uses_the_correct_setting(monkeypatch, purpose, light, expected):
    values = {**db.SETTING_DEFAULTS, 'openrouter_reasoning': 'low', 'openrouter_workshop_reasoning': 'xhigh'}
    monkeypatch.setattr(ai_provider, 'settings', lambda: values)
    monkeypatch.setattr(ai_provider, 'configuration', lambda *a, **kw: ('openrouter', 'test', 'key'))
    seen = []
    monkeypatch.setattr(ai_provider, '_call', lambda *args: seen.append(args[-1]['_reasoning_effort']) or 'ok')
    assert ai_provider.generate_content(model='test', config=SimpleNamespace(), contents='Preguntas',
                                        light=light, reasoning_purpose=purpose) == 'ok'
    assert seen == [expected]


def test_workshop_reasoning_is_preserved_when_falling_back_to_openrouter(monkeypatch):
    values = {**db.SETTING_DEFAULTS, 'ai_provider': 'gemini', 'fallback_enabled': '1',
              'openrouter_reasoning': 'low', 'openrouter_workshop_reasoning': 'high'}
    monkeypatch.setattr(ai_provider, 'settings', lambda: values)
    monkeypatch.setattr(ai_provider, 'configuration', lambda *a, **kw: ('gemini', 'test', 'key'))
    monkeypatch.setattr(ai_provider, '_fallback_configuration', lambda *a: ('openrouter', 'test', 'key'))
    seen = []
    def call(provider, model, key, config, contents, values):
        seen.append((provider, values['_reasoning_effort']))
        if provider == 'gemini':
            raise IA.GenerationError('unavailable', 'No disponible')
        return 'ok'
    monkeypatch.setattr(ai_provider, '_call', call)
    assert ai_provider.generate_content(model='test', config=SimpleNamespace(), contents='Preguntas',
                                        reasoning_purpose='workshop') == 'ok'
    assert seen == [('gemini', 'high'), ('openrouter', 'high')]
