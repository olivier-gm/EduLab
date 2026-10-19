"""Regression: math stays complete, editable and outside the table of contents."""
import pytest
from docx import Document
from docx.oxml.ns import qn

import IA
from algorythms import Document_process
import math_format


@pytest.mark.parametrize('formula', [
    r'\(\neg(p\land q)\Leftrightarrow(\neg p)\lor(\neg q)\)',
    r'\(x\neq1,\quad x\in\mathbb{R}\)',
    r'\[\left(\frac{x^2-1}{x-1}\right)=x+1\]',
    r'$\sqrt{x^2+1}$', r'$$\int_0^1 x^2\,dx=\frac{1}{3}$$',
])
def test_cleanup_preserves_formulas_and_literal_paragraph_breaks(formula):
    assert IA._clean('output: Antes.\\n\\n' + formula + '\\nDespués.') == (
        'Antes.\n\n' + formula + '\nDespués.')


def test_unwrapped_commands_are_not_newline_escapes():
    assert IA._clean(r'\neg p, x\neq1, \right), \rho') == r'\neg p, x\neq1, \right), \rho'


@pytest.mark.parametrize('formula,structure', [
    (r'\(\frac{x^2-1}{x-1}\)', 'f'),
    (r'\(\sqrt[3]{x+1}\)', 'rad'),
    (r'\(\sqrt{4-x^2}\)', 'rad'),
    (r'\(x_i^2\)', 'sSubSup'),
    (r'\[\sum_{i=1}^{n}i=\frac{n(n+1)}{2}\]', 'nary'),
    (r'\[\int_0^1 x^2\,dx=\frac{1}{3}\]', 'nary'),
    (r'\[\begin{cases}x & x\geq0\\-x & x<0\end{cases}\]', 'm'),
    (r'\[\begin{aligned}x&=2+3\\&=5\end{aligned}\]', 'm'),
    (r'\[\begin{pmatrix}1&2\\3&4\end{pmatrix}\]', 'm'),
])
def test_real_office_equations_survive_save_and_reload(tmp_path, formula, structure):
    doc = Document()
    Document_process.parrafos('Explicación.\n' + formula, doc, 'Matemáticas', False)
    target = tmp_path / 'math.docx'
    doc.save(target)
    loaded = Document(target)
    assert loaded._element.findall('.//' + qn('m:oMath'))
    assert loaded._element.findall('.//' + qn('m:' + structure))
    assert [p.text for p in loaded.paragraphs if p.style.name.startswith('Heading')] == ['Matemáticas']
    assert '\\' not in ''.join(loaded._element.itertext())
    assert loaded.paragraphs[-1].paragraph_format.line_spacing == 1.4


def test_de_morgan_does_not_split_or_enter_the_index():
    formula = r'\(\neg(p\land q)\Leftrightarrow(\neg p)\lor(\neg q)\)'
    doc = Document()
    Document_process.parrafos(IA._clean('De Morgan: ' + formula + '.'), doc, 'Lógica', False)
    assert len(doc.paragraphs) == 2
    assert len(doc._element.findall('.//' + qn('m:oMath'))) == 1
    text = ''.join(doc._element.itertext())
    assert '¬' in text and '∧' in text and '∨' in text
    assert 'Eg' not in text


def test_multiline_formula_is_one_block():
    formula = '\\[\\begin{aligned}\nx&=2+3\\\\\n&=5\n\\end{aligned}\\]'
    assert Document_process._split_blocks('Antes.\n' + formula + '\nDespués.') == [
        'Antes.', formula, 'Después.']


def test_markdown_and_multiplication_survive_next_to_equations():
    doc = Document()
    p = doc.add_paragraph()
    Document_process._add_inline(p, r'**Resultado: \(a*b=c\)** y *explicación*.')
    assert any(r.bold and 'Resultado' in r.text for r in p.runs)
    assert any(r.italic and 'explicación' in r.text for r in p.runs)
    assert '*' in ''.join(p._p.itertext())
    assert len(p._p.findall(qn('m:oMath'))) == 1


def test_ordinary_money_and_calculations_stay_text():
    text = 'Precio $5 y $10. Juego mínimo = 80,032 − 80,030 = 0,002 mm; 2*3 = 6.'
    assert not math_format.has_math(text)
    p = Document().add_paragraph()
    Document_process._add_inline(p, text)
    assert p.text == text
    assert not p._p.findall(qn('m:oMath'))


def test_conversion_failure_preserves_every_operand(monkeypatch):
    def fail(*args):
        raise ValueError('unsupported')
    monkeypatch.setattr(math_format, '_equation_xml', fail)
    p = Document().add_paragraph()
    Document_process._add_inline(p, r'Antes \(x^2+2x+1=0\) después.')
    assert p.text == r'Antes \(x^2+2x+1=0\) después.'


@pytest.mark.parametrize('text', [r'\(x^2\)', r'\[\sqrt{x}\]', 'x^2 + 2*x = 5', 'EI = 0'])
def test_formulas_are_not_heading_candidates(text):
    assert not Document_process._is_subtitle(text)


@pytest.mark.parametrize('detect', [True, False])
def test_apa_first_line_indent_with_equation_and_list_exceptions(detect):
    doc = Document()
    Document_process.parrafos(
        'Primer párrafo con contenido suficiente para ser una oración.\n'
        'Segundo párrafo con \\(x^2\\) dentro del texto.\n'
        '\\[\\frac{1}{2}+\\frac{1}{2}=1\\]\n'
        '- Un elemento de lista.', doc, 'Desarrollo', False, detect)
    assert doc.paragraphs[0].paragraph_format.first_line_indent is None
    assert doc.paragraphs[1].paragraph_format.first_line_indent.cm == pytest.approx(1.27)
    assert doc.paragraphs[2].paragraph_format.first_line_indent.cm == pytest.approx(1.27)
    assert doc.paragraphs[3].paragraph_format.first_line_indent == 0
    if detect:
        assert doc.paragraphs[4].paragraph_format.first_line_indent.cm < 0


def test_apa_references_use_hanging_indent():
    doc = Document()
    Document_process.parrafos('Autor (2026). Una referencia.', doc, 'Bibliografía', False, False)
    reference = doc.paragraphs[1]
    assert reference.paragraph_format.left_indent.cm == pytest.approx(1.27)
    assert reference.paragraph_format.first_line_indent.cm == pytest.approx(-1.27)


def test_square_root_has_required_degree_slot_for_word_and_libreoffice():
    p = Document().add_paragraph()
    Document_process._add_inline(p, r'\(\sqrt{4-x^2}\)')
    radical = p._p.find('.//' + qn('m:rad'))
    assert radical.find(qn('m:deg')) is not None
    assert radical.find('.//' + qn('m:degHide')).get(qn('m:val')) == '1'
    assert '4' in ''.join(radical.find(qn('m:e')).itertext())


def test_repeated_equations_are_not_limited_or_moved_by_the_cache():
    doc = Document()
    formulas = '\n'.join(r'\[x^2+1=2\]' for _ in range(80))
    Document_process.parrafos(formulas, doc, 'Cálculos', False)
    assert len(doc._element.findall('.//' + qn('m:oMath'))) == 80
    assert len(doc.paragraphs) == 81


def test_requested_heading_can_contain_an_editable_formula():
    doc = Document()
    heading = r'Función \(f(x)=x^2\)'
    Document_process.parrafos(heading + '\nContenido del apartado.', doc, 'Estudio', False,
                             known_subtitles=[heading])
    assert doc.paragraphs[1].style.name == 'Heading 2'
    assert len(doc.paragraphs[1]._p.findall(qn('m:oMath'))) == 1
