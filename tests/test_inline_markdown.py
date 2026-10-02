# -*- coding: utf-8 -*-
"""Marcas de énfasis de la IA (**negrita**, *cursiva*): se convierten en formato
real de Word y nunca quedan asteriscos sueltos en el documento."""
import pytest
from docx import Document

from algorythms import Document_process

seg = Document_process._inline_segments


def formatted(paragraph):
    """[(texto, negrita, cursiva)] de los runs de un párrafo, sin runs vacíos."""
    return [(r.text, bool(r.bold), bool(r.italic)) for r in paragraph.runs if r.text]


# ── Interpretación ────────────────────────────────────────────────────

def test_cursiva_con_un_asterisco_como_en_el_puma():
    assert seg('conocido como *Puma concolor*, es un felino') == [
        ('conocido como ', False, False), ('Puma concolor', False, True), (', es un felino', False, False)]


def test_cursiva_entre_parentesis():
    assert seg('del puma (*Puma concolor*) nos permite') == [
        ('del puma (', False, False), ('Puma concolor', False, True), (') nos permite', False, False)]


def test_negrita_con_dos_asteriscos():
    assert seg('esto es **muy importante** hoy') == [
        ('esto es ', False, False), ('muy importante', True, False), (' hoy', False, False)]


def test_negrita_cursiva_con_tres():
    assert seg('***clave***') == [('clave', True, True)]


def test_cursiva_dentro_de_negrita():
    assert seg('**el *Puma concolor* vive**') == [
        ('el ', True, False), ('Puma concolor', True, True), (' vive', True, False)]


def test_varias_marcas_en_el_mismo_parrafo():
    assert [(t, b, i) for t, b, i in seg('*a* y **b** y *c*')] == [
        ('a', False, True), (' y ', False, False), ('b', True, False), (' y ', False, False), ('c', False, True)]


@pytest.mark.parametrize('text', ['2*3*4 = 24', 'precio 5*2', 'a * b * c', 'E = m * c', '3 * 4 = 12'])
def test_multiplicaciones_y_asteriscos_con_espacios_no_son_marcas(text):
    assert seg(text) == [(text, False, False)]


@pytest.mark.parametrize('text,expected', [
    ('texto con * suelto', 'texto con  suelto'),
    ('abre *sin cerrar', 'abre sin cerrar'),
    ('cierra sin abrir*', 'cierra sin abrir'),
    ('doble ** vacío', 'doble  vacío'),
])
def test_asteriscos_sueltos_se_eliminan_en_texto_de_la_ia(text, expected):
    assert ''.join(t for t, _, _ in seg(text)) == expected


def test_en_texto_manual_los_asteriscos_sueltos_se_conservan():
    assert seg('Nota importante*', strip_stray=False) == [('Nota importante*', False, False)]


def test_strip_inline_markers():
    assert Document_process._strip_inline_markers('**Origen** del *Puma*') == 'Origen del Puma'


# ── En el documento ───────────────────────────────────────────────────

def build(body, detect=True):
    doc = Document()
    Document_process.parrafos(body, doc, 'Titulo', False, detect)
    return doc


def test_el_documento_no_contiene_asteriscos_y_la_cursiva_es_real():
    frase = ('El Puma, conocido científicamente como *Puma concolor*, es un mamífero '
             'carnívoro perteneciente a la familia Felidae.')
    doc = build(frase)
    paragraph = doc.paragraphs[1]
    assert '*' not in paragraph.text
    assert ('Puma concolor', False, True) in formatted(paragraph)
    assert all(r.font.name == 'Arial' for r in paragraph.runs)


def test_la_frase_de_la_conclusion_del_puma():
    doc = build('Para concluir, el estudio del puma (*Puma concolor*) nos permite comprender la adaptabilidad.')
    assert '*' not in doc.paragraphs[1].text
    assert ('Puma concolor', False, True) in formatted(doc.paragraphs[1])


def test_los_subtitulos_no_muestran_marcas():
    doc = build('**Origen y Fundamentos**\n\nUn párrafo normal con **negrita**.')
    heading, body = doc.paragraphs[1], doc.paragraphs[2]
    assert heading.text == 'Origen y Fundamentos' and heading.style.name == 'Heading 2'
    assert ('negrita', True, False) in formatted(body)


def test_las_vinetas_se_convierten_en_lista():
    doc = build('Texto introductorio.\n\n* Primer punto\n* Segundo con **énfasis**')
    texts = [p.text for p in doc.paragraphs[2:]]
    assert texts[0].startswith('• Primer punto') and texts[1].startswith('• Segundo')
    assert all('*' not in t for t in texts)
    assert doc.paragraphs[2].style.name != 'Heading 2'


def test_una_vineta_corta_no_se_toma_por_subtitulo():
    doc = build('* Punto breve\n* Otro punto')
    assert all(p.style.name != 'Heading 2' for p in doc.paragraphs[1:])


def test_el_texto_manual_aplica_las_parejas_pero_conserva_asteriscos_sueltos():
    doc = build('Esto es *cursiva* y una nota*.', detect=False)
    assert doc.paragraphs[1].text == 'Esto es cursiva y una nota*.'
    assert ('cursiva', False, True) in formatted(doc.paragraphs[1])


def test_el_glosario_aplica_las_marcas():
    doc = Document()
    Document_process.add_glossary(doc, [{'term': '**Puma**', 'definition': 'Felino *Puma concolor* americano.'}])
    paragraph = next(p for p in doc.paragraphs if p.text.startswith('Puma'))
    assert paragraph.text == 'Puma: Felino Puma concolor americano.'
    assert ('Puma concolor', False, True) in formatted(paragraph)
