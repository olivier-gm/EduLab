from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.shared import Cm
from PIL import Image

from algorythms import Document_process


@pytest.mark.parametrize('extension', ['png', 'webp', 'svg'])
def test_logo_is_embedded_in_word_with_transparency(tmp_path, extension):
    path = tmp_path / f'universidad_de_prueba.{extension}'
    if extension == 'svg':
        path.write_text(
            '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 40 40">'
            '<rect x="10" y="10" width="20" height="20" fill="#00cccc"/>'
            '</svg>', encoding='utf-8')
    else:
        picture = Image.new('RGBA', (40, 40))
        picture.paste((0, 204, 204, 255), (10, 10, 30, 30))
        picture.save(path, lossless=True)
    document = Document()
    document.add_paragraph('Portada')
    assert Document_process.insert_logo(document, 'Universidad de Prúeba', str(tmp_path))
    assert len(document.inline_shapes) == 1
    assert document.inline_shapes[0].width == Cm(3)
    output = BytesIO()
    document.save(output)
    with ZipFile(output) as archive:
        media = [name for name in archive.namelist() if name.startswith('word/media/')]
        assert len(media) == 1
        with Image.open(BytesIO(archive.read(media[0]))) as embedded:
            assert embedded.format == 'PNG'
            assert embedded.convert('RGBA').getpixel((0, 0))[3] == 0
    assert document.paragraphs[1].text == 'Portada'


def test_invalid_svg_does_not_change_cover(tmp_path):
    (tmp_path / 'universidad.svg').write_text('invalid svg', encoding='utf-8')
    document = Document()
    document.add_paragraph('Portada')
    assert not Document_process.insert_logo(document, 'Universidad', str(tmp_path))
    assert [paragraph.text for paragraph in document.paragraphs] == ['Portada']
    assert not document.inline_shapes


# ── Mismo alto para todos los logos y posición bajo el membrete ──────

def _logo(tmp_path, width, height, name='universidad_de_prueba.png'):
    picture = Image.new('RGBA', (width, height), (0, 0, 0, 0))
    picture.paste((10, 20, 30, 255), (1, 1, width - 1, height - 1))
    picture.save(tmp_path / name)


@pytest.mark.parametrize('size', [(200, 280), (1000, 1000), (500, 200)])
def test_todos_los_logos_miden_lo_mismo_de_alto(tmp_path, size):
    _logo(tmp_path, *size)
    document = Document()
    document.add_paragraph('Portada')
    assert Document_process.insert_logo(document, 'Universidad de Prueba', str(tmp_path))
    shape = document.inline_shapes[0]
    assert abs(shape.height - Cm(Document_process.LOGO_HEIGHT_CM)) <= Cm(0.01)
    assert shape.width <= Cm(Document_process.LOGO_MAX_WIDTH_CM)
    assert abs(shape.width / shape.height - size[0] / size[1]) < 0.01          # sin deformarlo


def test_un_logo_muy_ancho_se_limita_de_ancho_sin_deformarse(tmp_path):
    _logo(tmp_path, 1500, 200)
    document = Document()
    document.add_paragraph('Portada')
    Document_process.insert_logo(document, 'Universidad de Prueba', str(tmp_path))
    shape = document.inline_shapes[0]
    assert shape.width == Cm(Document_process.LOGO_MAX_WIDTH_CM)
    assert shape.height < Cm(Document_process.LOGO_HEIGHT_CM)
    assert abs(document._logo_height_pt - shape.height / 12700) < 0.01


def test_el_logo_va_debajo_del_membrete_y_antes_del_titulo(tmp_path):
    _logo(tmp_path, 100, 100)
    document = Document()
    for text in ('REPÚBLICA', 'MINISTERIO', 'UNIVERSIDAD', '[carrera]', '', '', '[title]', 'DOCENTE'):
        document.add_paragraph(text)
    assert Document_process.insert_logo(document, 'Universidad de Prueba', str(tmp_path))
    texts = [('LOGO' if Document_process._has_drawing(p) else p.text) for p in document.paragraphs]
    assert texts == ['REPÚBLICA', 'MINISTERIO', 'UNIVERSIDAD', '[carrera]', 'LOGO', '', '', '[title]', 'DOCENTE']


# ── tools/fix_logos.py ───────────────────────────────────────────────

def test_fix_logos_recorta_margenes_y_es_idempotente(tmp_path):
    import importlib.util
    import os
    spec = importlib.util.spec_from_file_location(
        'fix_logos', os.path.join(os.path.dirname(__file__), '..', 'tools', 'fix_logos.py'))
    fix = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(fix)

    path = tmp_path / 'universidad_de_prueba.png'
    picture = Image.new('RGBA', (400, 900), (0, 0, 0, 0))
    picture.paste((200, 0, 0, 255), (100, 100, 300, 700))          # 200x600 visible
    picture.save(path)
    before, after, notes, changed = fix.fix_raster(str(path), check_only=True)
    assert changed and after == (200, 600) and Image.open(path).size == (400, 900)   # --check no escribe
    fix.fix_raster(str(path), check_only=False)
    assert Image.open(path).size == (200, 600)
    assert fix.fix_raster(str(path), check_only=False)[3] is False                 # segunda vez: nada que hacer
