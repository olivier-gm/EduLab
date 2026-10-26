import pytest
from docx import Document

from algorythms import Document_process
from form_processor import FormProcessor


@pytest.mark.parametrize('route', ['uni', 'bach'])
@pytest.mark.parametrize('font_name', ['Arial', 'Times New Roman'])
@pytest.mark.parametrize('count', [1, 8])
def test_long_names_keep_complete_identity_inside_two_columns(tmp_path, monkeypatch, route, font_name, count):
    data = {'u': 'UNERG', 'title': 'Ingeniería Económica', 'teacher': 'Gabriela Torres',
            'city': 'Caracas', 'date': '25/10/2026', 'seccion': '3',
            'gblock-template-canvas-integrantes': str(count)}
    for index in range(1, count + 1):
        data[f'input{index}'] = 'Jefferso Enrique Santamaria Santamaria'
        data[f'id{index}'] = str(28408420 + index)
    processor = FormProcessor(data, route)
    processor.process()
    replacements, title = processor.generate_replacements()
    monkeypatch.setattr(Document_process, 'convert', lambda *a: None)
    output = tmp_path / 'cover.docx'
    template = 'input/plantilla_bach.docx' if route == 'bach' else 'input/plantilla.docx'
    Document_process.fill_placeholders(str(output), template, '', replacements, '', '', '', title, route,
                                       font_name=font_name)
    document = Document(output)
    assert len(document.tables) == 1
    table = document.tables[0]
    assert 'GABRIELA TORRES' in table.cell(0, 0).text
    assert 'SECCION: "3"' in table.cell(0, 0).text
    students = table.cell(0, 1).paragraphs[1:]
    assert len(students) == count
    for index, paragraph in enumerate(students, 1):
        assert paragraph.text == replacements[f'[{index}]'] + '\n' + replacements[f'[id{index}]'].strip()
        assert all(run.font.name == font_name for run in paragraph.runs)
    assert table.rows[0]._tr.xpath('./w:trPr/w:cantSplit')
    assert not any('[id' in p.text for row in table.rows for cell in row.cells for p in cell.paragraphs)
    assert any('25 DE OCTUBRE DE 2026' in p.text for p in document.paragraphs)


def test_short_names_keep_existing_cover_layout(tmp_path, monkeypatch):
    processor = FormProcessor({'title': 'Tema de prueba', 'input1': 'Ana Pérez', 'id1': '12345678',
                               'gblock-template-canvas-integrantes': '1'}, 'uni')
    processor.process()
    replacements, title = processor.generate_replacements()
    monkeypatch.setattr(Document_process, 'convert', lambda *a: None)
    output = tmp_path / 'cover.docx'
    Document_process.fill_placeholders(str(output), 'input/plantilla.docx', '', replacements,
                                       '', '', '', title, 'uni')
    document = Document(output)
    assert not document.tables
    assert any('ANA PÉREZ C.I- 12345678' in p.text for p in document.paragraphs)
