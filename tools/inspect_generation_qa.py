"""Fixtures y revisión de Word/PDF con las dependencias del runtime de Codex."""
import argparse
import io
import json
from pathlib import Path
import re
import unicodedata
import zipfile

from docx import Document
from PIL import Image, ImageDraw, ImageFont
import pypdfium2 as pdfium
from pypdf import PdfReader

TERMS = ['Átomo', 'Célula', 'ADN', 'ARN', 'Ósmosis', 'Difusión', 'Enzima', 'Mitosis', 'Meiosis', 'Tejido']


def write(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def fixtures(directory):
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    directory.mkdir(parents=True, exist_ok=True)
    header = ['UNIVERSIDAD DE PRUEBA', 'Facultad de Ciencias', 'Biología celular',
              'Asignación de glosario', 'Docente: María Pérez', 'Estudiante: Ana Gómez',
              'Entrega: 5 de octubre de 2026']
    instruction = 'Investiga únicamente estos diez términos. No agregues introducción ni conclusión.'
    lines = header + ['', instruction, ''] + [f'{i+1}. {term}' for i, term in enumerate(TERMS)]
    (directory / 'asignacion.txt').write_text('\n'.join(lines), encoding='utf-8')
    document = Document()
    for line in header:
        document.add_paragraph(line)
    document.add_page_break()
    document.add_paragraph(instruction)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = 'N.º'
    table.rows[0].cells[1].text = 'Término'
    for index, term in enumerate(TERMS, 1):
        row = table.add_row()
        row.cells[0].text = str(index)
        row.cells[1].text = term
    document.save(directory / 'asignacion.docx')
    font_path = Path('C:/Windows/Fonts/arial.ttf')
    pdfmetrics.registerFont(TTFont('QAArial', str(font_path)))
    target = directory / 'asignacion.pdf'
    pdf = canvas.Canvas(str(target), pagesize=(612, 792))
    pdf.setFont('QAArial', 16)
    for index, line in enumerate(header):
        pdf.drawString(60, 720 - index * 40, line)
    pdf.showPage()
    pdf.setFont('QAArial', 12)
    pdf.drawString(50, 740, instruction)
    for index, term in enumerate(TERMS, 1):
        pdf.drawString(70, 700 - index * 35, f'{index}. {term}')
    pdf.save()
    loaded = pdfium.PdfDocument(str(target))
    pages = [page.render(scale=1.5).to_pil().convert('RGB') for page in loaded]
    image = Image.new('RGB', (pages[0].width, sum(page.height for page in pages)), 'white')
    y = 0
    for page in pages:
        image.paste(page, (0, y))
        y += page.height
    for suffix in ('png', 'jpg', 'webp'):
        image.save(directory / f'asignacion.{suffix}')
    scanned = canvas.Canvas(str(directory / 'asignacion.scan.pdf'), pagesize=(612, 792))
    for page in pages:
        scanned.drawImage(ImageReader(page), 0, 0, width=612, height=792)
        scanned.showPage()
    scanned.save()
    loaded.close()
    assignment = Image.new('RGB', (1200, 1000), 'white')
    draw = ImageDraw.Draw(assignment)
    font = ImageFont.truetype(str(font_path), 30)
    report_lines = header + ['', 'Elabora un informe sobre el ciclo del agua.',
        'Incluye estos apartados:', '1. Evaporación', '2. Condensación', '3. Precipitación']
    for index, line in enumerate(report_lines):
        draw.text((45, 35 + index * 65), line, font=font, fill='black')
    for suffix in ('png', 'jpg', 'webp'):
        assignment.save(directory / f'informe.{suffix}')


def key(text):
    text = text.casefold().replace('ñ', '\uffff')
    return ''.join(c for c in unicodedata.normalize('NFD', text) if not unicodedata.combining(c)).replace('\uffff', 'n{')


def inspect(directory):
    request = json.loads((directory / 'solicitud.json').read_text(encoding='utf-8'))
    response = json.loads((directory / 'respuestas.json').read_text(encoding='utf-8'))
    form = request['form']
    content = response.get('content', {})
    document = Document(directory / 'documento.docx')
    text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    headings = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.style.name.startswith('Heading')]
    issues = []
    glossary = form.get('document_kind') == 'glossary'
    for label, option in [('Introducción', 'incluir_introduccion'), ('Conclusión', 'incluir_conclusion')]:
        if not glossary and (label in headings) != (option in form):
            issues.append(f'Sección {label}: no coincide con la opción solicitada.')
    if not glossary and ('Bibliografía' in headings) != ('incluir_bibliografia' in form):
        issues.append('La bibliografía no coincide con la opción solicitada.')
    if not glossary:
        for subtitle in content.get('subtitles', []):
            if key(subtitle) not in set(map(key, headings)):
                issues.append(f'Falta el subtítulo solicitado: {subtitle}.')
    if glossary:
        if any(label in headings for label in ('Introducción', 'Conclusión', 'Bibliografía')):
            issues.append('El glosario contiene una sección final o introductoria no permitida.')
        entries = content.get('glossary_entries') or []
        terms = [entry['term'] for entry in entries]
        if len(terms) != request.get('expected_count'):
            issues.append('Cantidad de términos incorrecta.')
        if terms != sorted(terms, key=key) or len(set(map(key, terms))) != len(terms):
            issues.append('Los términos no son únicos o no están ordenados alfabéticamente.')
        if form.get('glossary_source') == 'list':
            expected = (json.loads((directory / 'extraccion.json').read_text(encoding='utf-8'))['terms']
                        if request.get('upload') else form['glossary_terms'].splitlines())
            if set(map(key, expected)) != set(map(key, terms)):
                issues.append('Se omitieron o cambiaron términos predefinidos.')
        if any(not 1 <= len(entry['definition'].split()) <= 40 for entry in entries):
            issues.append('Definición vacía o superior a 40 palabras.')
        if 'incluir_bibliografia' in form and any(not entry.get('reference') for entry in entries):
            issues.append('Falta la fuente en uno o más términos.')
        normalized_word = ' '.join(text.split())
        if any(f'{entry["term"]}: {entry["definition"]}' not in normalized_word for entry in entries):
            issues.append('No se conservaron todas las definiciones completas en el Word.')
    if re.search(r'\[(?:title|date|u|name|input\d+)\]|\\n|^output:', text, re.M | re.I):
        issues.append('Marcadores de plantilla o escapes visibles en el Word.')
    if content.get('body') and len(content['body'].split()) < 50 and form.get('global-mode') == 'ia' and not glossary:
        issues.append('Desarrollo de IA demasiado breve.')
    result = {'issues': issues, 'headings': headings, 'inline_images': len(document.inline_shapes), 'characters_word': len(text)}
    result['fonts_word'] = sorted({run.font.name for paragraph in document.paragraphs for run in paragraph.runs if run.font.name and run.text.strip()})
    if form.get('fuente') == 'tnr' and any(font != 'Times New Roman' for font in result['fonts_word']):
        issues.append('La fuente del Word no coincide con Times New Roman.')
    pdf_path = directory / 'documento.pdf'
    if not pdf_path.exists():
        issues.append('Falta el PDF.')
    else:
        pdf = pdfium.PdfDocument(str(pdf_path))
        page_texts = [page.get_textpage().get_text_range() for page in pdf]
        result['pages'] = len(pdf)
        result['characters_pdf'] = sum(map(len, page_texts))
        result['blank_pages'] = [i+1 for i, value in enumerate(page_texts) if not value.strip()]
        if result['blank_pages']:
            issues.append(f'Páginas vacías: {result["blank_pages"]}')
        first = ' '.join(page_texts[0].split())
        section_start = 'Glosario' if glossary else 'Índice'
        if (glossary or any(content.get(key) for key in ('body', 'introduction', 'conclusion', 'bibliography'))) and len(page_texts) > 1:
            section_pages = [i for i, value in enumerate(page_texts) if re.search(rf'(?m)^{section_start}\s*$', value)]
            if section_pages and section_pages[0] > 1:
                issues.append('La portada ocupa más de una página; revisar dimensiones del logo y campos de portada.')
        university_header = 'EDUCACIÓN UNIVERSITARIA' in first
        if form.get('instituto') == 'bachiller' and university_header:
            issues.append('Bachillerato usa el encabezado de educación universitaria.')
        if form.get('instituto') == 'universidad' and not university_header:
            issues.append('Universidad no usa el encabezado de educación universitaria.')
        if 'Bibliografía' in headings:
            bibliography_pages = [i for i, value in enumerate(page_texts) if re.search(r'(?m)^Bibliografía\s*$', value)]
            conclusion_pages = [i for i, value in enumerate(page_texts) if 'CONCLUSION_MANUAL' in value]
            if bibliography_pages and conclusion_pages and min(bibliography_pages) <= max(conclusion_pages):
                issues.append('Bibliografía no empieza en una página posterior a la conclusión.')
        if not glossary and content.get('body') and len(page_texts) > 1:
            if 'No se encontraron' in page_texts[1] or ('Índice' in page_texts[1] and len(page_texts[1].split()) < 5):
                issues.append('Tabla de contenido vacía o sin actualizar.')
        page_dir = directory / 'paginas'
        page_dir.mkdir(exist_ok=True)
        # Una repetición puede tener menos páginas; retirar solo nuestras vistas
        # numeradas dentro del directorio de QA, nunca imágenes del proyecto.
        qa_root = Path(__file__).resolve().parents[1] / 'qa_generaciones'
        if directory.resolve().is_relative_to(qa_root.resolve()):
            for previous in list(page_dir.glob('*.png')) + list(directory.glob('vista_*.png')):
                if re.fullmatch(r'(?:\d{3}|vista_\d+)\.png', previous.name):
                    previous.unlink()
        thumbnails = []
        font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 20)
        for index, page in enumerate(pdf):
            image = page.render(scale=1.5).to_pil().convert('RGB')
            image.save(page_dir / f'{index+1:03}.png')
            image.thumbnail((450, 600))
            thumbnail = Image.new('RGB', (470, 640), '#dedee8')
            thumbnail.paste(image, ((470-image.width)//2, 30))
            ImageDraw.Draw(thumbnail).text((12, 6), f'{request["id"]} · página {index+1}', font=font, fill='black')
            thumbnails.append(thumbnail)
        # Hojas de revisión de 4 páginas, además de las páginas a escala mayor.
        for offset in range(0, len(thumbnails), 4):
            sheet = Image.new('RGB', (940, 1280), 'white')
            for index, image in enumerate(thumbnails[offset:offset+4]):
                sheet.paste(image, ((index % 2)*470, (index // 2)*640))
            sheet.save(directory / f'vista_{offset//4+1:02}.png')
        pdf.close()
        reader = PdfReader(str(pdf_path))
        result['pdf_links'] = sum(1 for page in reader.pages for annotation in page.get('/Annots', [])
                                  if annotation.get_object().get('/Subtype') == '/Link')
        external_urls = set()
        for page in reader.pages:
            for annotation in page.get('/Annots', []):
                action = annotation.get_object().get('/A')
                if action and action.get_object().get('/URI'):
                    external_urls.add(str(action.get_object()['/URI']))
        result['pdf_external_urls'] = sorted(external_urls)
    (directory / 'texto_word.txt').write_text(text, encoding='utf-8')
    if pdf_path.exists():
        (directory / 'texto_pdf.txt').write_text('\n\n'.join(f'PÁGINA {i+1}\n{value}' for i, value in enumerate(page_texts)), encoding='utf-8')
    write(directory / 'inspeccion.json', result)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--fixtures', type=Path)
    parser.add_argument('--case', type=Path)
    args = parser.parse_args()
    if args.fixtures:
        fixtures(args.fixtures)
    if args.case:
        inspect(args.case)
